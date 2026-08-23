import sys, os, re, subprocess, markdown
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
DOCS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DIAGRAMS_DIR = os.path.join(DOCS_DIR, "diagrams")

CORAL = "FF6B4A"
CORAL_DARK = "F0512E"
NAVY = "1B2A4A"
NAVY2 = "243B60"
YELLOW = "FFC857"
INK = "182338"
BG = "FFF8F5"
MUTED = "64748B"
LINE = "E5E7EB"

PAGE_CSS = f"""
@page {{ size: A4; margin: 20mm 16mm 20mm 16mm; }}
* {{ box-sizing: border-box; }}
body {{ font-family: Arial, "Segoe UI", sans-serif; color: #{INK}; font-size: 11.5pt; line-height: 1.55; }}
.cover {{
  background: linear-gradient(135deg, #{NAVY}, #{NAVY2});
  color: #fff; border-radius: 18px; padding: 70px 56px; margin-bottom: 26px;
  page-break-after: always;
}}
.cover .kicker {{ color: #{YELLOW}; font-weight: 800; letter-spacing: .12em; text-transform: uppercase; font-size: 12px; margin-bottom: 14px; }}
.cover h1 {{ color: #fff; font-size: 40px; margin: 0 0 14px; font-weight: 800; }}
.cover .sub {{ color: #D7DEEC; font-size: 16px; max-width: 640px; }}
.cover .date {{ color: #94A3B8; font-size: 12px; margin-top: 40px; }}
h1 {{ color: #{NAVY}; font-size: 21pt; border-bottom: 3px solid #{CORAL}; padding-bottom: 6px; margin-top: 30px; page-break-after: avoid; }}
h2 {{ color: #{NAVY}; font-size: 15.5pt; margin-top: 24px; page-break-after: avoid; }}
h3 {{ color: #{CORAL_DARK}; font-size: 12.5pt; margin-top: 16px; page-break-after: avoid; }}
h4 {{ color: #{INK}; font-size: 11.5pt; margin-top: 12px; page-break-after: avoid; }}
p {{ margin: 8px 0; }}
a {{ color: #{CORAL_DARK}; }}
ul, ol {{ margin: 6px 0; padding-left: 22px; }}
li {{ margin-bottom: 4px; }}
table {{ border-collapse: collapse; width: 100%; margin: 14px 0; font-size: 10pt; }}
th {{ background: #{NAVY}; color: #fff; text-align: left; padding: 7px 10px; }}
td {{ border-bottom: 1px solid #{LINE}; padding: 6px 10px; vertical-align: top; }}
tr:nth-child(even) td {{ background: #FAFAFA; }}
code {{ background: #F1F5F9; padding: 1px 5px; border-radius: 4px; font-size: 10pt; }}
pre {{ background: #182338; color: #E2E8F0; padding: 14px 16px; border-radius: 10px; overflow-x: auto; font-size: 9.5pt; }}
pre code {{ background: none; color: inherit; }}
blockquote {{ border-left: 4px solid #{YELLOW}; background: #FFF8F5; margin: 12px 0; padding: 8px 16px; color: #{MUTED}; }}
img {{ max-width: 100%; display: block; margin: 14px auto; border: 1px solid #{LINE}; border-radius: 10px; }}
hr {{ border: none; border-top: 1px solid #{LINE}; margin: 26px 0; }}
.figcap {{ text-align: center; font-size: 9.5pt; color: #{MUTED}; margin-top: -8px; margin-bottom: 18px; }}
"""

def md_to_body_html(md_path):
    text = open(md_path, encoding="utf-8").read()
    html = markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists", "toc"])
    md_dir = os.path.dirname(os.path.abspath(md_path))
    def rewrite(m):
        src = m.group(1)
        if src.startswith("http://") or src.startswith("https://") or src.startswith("file://"):
            return m.group(0)
        abs_path = os.path.normpath(os.path.join(md_dir, src))
        url = "file:///" + abs_path.replace("\\", "/")
        return m.group(0).replace(src, url)
    html = re.sub(r'<img[^>]*src="([^"]+)"', rewrite, html)
    return html

def make_cover_html(title, subtitle, badges, date_str):
    badge_html = "".join(f'<span style="background:rgba(255,255,255,.14);border:1px solid rgba(255,255,255,.35);border-radius:999px;padding:6px 14px;font-size:11px;font-weight:700;margin-right:8px;display:inline-block;margin-top:8px;">{b}</span>' for b in badges)
    return f"""
    <div class="cover">
      <div class="kicker">Basera Platform Documentation</div>
      <h1>{title}</h1>
      <div class="sub">{subtitle}</div>
      <div>{badge_html}</div>
      <div class="date">{date_str}</div>
    </div>
    """

def build_html(md_path, out_html, title, subtitle, badges, date_str):
    body = md_to_body_html(md_path)
    cover = make_cover_html(title, subtitle, badges, date_str)
    full = f"""<!doctype html><html><head><meta charset="utf-8"><title>{title}</title>
    <style>{PAGE_CSS}</style></head><body>{cover}{body}</body></html>"""
    open(out_html, "w", encoding="utf-8").write(full)
    return out_html

def html_to_pdf(html_path, pdf_path):
    file_url = "file:///" + os.path.abspath(html_path).replace("\\", "/")
    cmd = [CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
           f"--print-to-pdf={os.path.abspath(pdf_path)}",
           "--print-to-pdf-no-header", file_url]
    r = subprocess.run(cmd, capture_output=True, text=True)
    return os.path.exists(pdf_path)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def add_hyperlink_run(paragraph, text):
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor.from_string(CORAL_DARK)
    return run

def add_runs_from_inline(paragraph, node):
    """Walk inline children of a tag, adding formatted runs to a docx paragraph."""
    for child in node.children:
        if isinstance(child, str):
            if child:
                paragraph.add_run(child)
        else:
            name = child.name
            txt = child.get_text()
            if name in ("strong", "b"):
                r = paragraph.add_run(txt); r.bold = True
            elif name in ("em", "i"):
                r = paragraph.add_run(txt); r.italic = True
            elif name == "code":
                r = paragraph.add_run(txt); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            elif name == "a":
                add_hyperlink_run(paragraph, txt)
            elif name == "br":
                paragraph.add_run().add_break()
            else:
                add_runs_from_inline(paragraph, child)

def add_image(doc, src):
    path = src
    if path.startswith("file:///"):
        path = path[8:]
    elif path.startswith("file://"):
        path = path[7:]
    path = path.replace("/", os.sep)
    if not os.path.isabs(path):
        path = os.path.normpath(os.path.join(DOCS_DIR, path))
    if os.path.exists(path):
        try:
            doc.add_picture(path, width=Inches(6.3))
            last = doc.paragraphs[-1]
            last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            doc.add_paragraph(f"[image: {src}]")

def html_to_docx(html_path, docx_path, title):
    html = open(html_path, encoding="utf-8").read()
    soup = BeautifulSoup(html, "html.parser")
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    # Cover
    cover = soup.find("div", class_="cover")
    if cover:
        h1 = cover.find("h1")
        sub = cover.find("div", class_="sub")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("BASERA")
        r.bold = True; r.font.size = Pt(34); r.font.color.rgb = RGBColor.from_string(NAVY)
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(h1.get_text() if h1 else title)
        r2.bold = True; r2.font.size = Pt(20); r2.font.color.rgb = RGBColor.from_string(CORAL_DARK)
        if sub:
            p3 = doc.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r3 = p3.add_run(sub.get_text())
            r3.font.size = Pt(12); r3.font.color.rgb = RGBColor.from_string(MUTED)
        doc.add_page_break()
        body_nodes = cover.find_all_next()
        top_level = [n for n in soup.body.children if getattr(n, "name", None) and n is not cover]
    else:
        top_level = [n for n in soup.body.children if getattr(n, "name", None)]

    for node in soup.body.children:
        name = getattr(node, "name", None)
        if not name or node is cover:
            continue
        if name in ("h1", "h2", "h3", "h4"):
            level = int(name[1])
            h = doc.add_heading(node.get_text(), level=level)
            for run in h.runs:
                run.font.color.rgb = RGBColor.from_string(NAVY if level <= 2 else CORAL_DARK)
        elif name == "p":
            img = node.find("img")
            if img:
                add_image(doc, img.get("src", ""))
                continue
            if not node.get_text(strip=True):
                continue
            p = doc.add_paragraph()
            add_runs_from_inline(p, node)
        elif name in ("ul", "ol"):
            for li in node.find_all("li", recursive=False):
                style_name = "List Bullet" if name == "ul" else "List Number"
                p = doc.add_paragraph(style=style_name)
                add_runs_from_inline(p, li)
        elif name == "table":
            rows = node.find_all("tr")
            if not rows:
                continue
            ncols = max(len(r.find_all(["td", "th"])) for r in rows)
            t = doc.add_table(rows=0, cols=ncols)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.style = "Light Grid Accent 1"
            for ri, row in enumerate(rows):
                cells = row.find_all(["td", "th"])
                tc = t.add_row().cells
                is_header = row.find("th") is not None
                for ci, cell in enumerate(cells):
                    if ci >= ncols:
                        break
                    tc[ci].text = ""
                    p = tc[ci].paragraphs[0]
                    add_runs_from_inline(p, cell)
                    if is_header:
                        set_cell_shading(tc[ci], NAVY)
                        for run in p.runs:
                            run.font.color.rgb = RGBColor.from_string("FFFFFF")
                            run.bold = True
        elif name == "blockquote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs_from_inline(p, node)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string(MUTED)
        elif name == "pre":
            p = doc.add_paragraph()
            r = p.add_run(node.get_text())
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif name == "hr":
            doc.add_paragraph("_" * 60)
        elif name == "img":
            add_image(doc, node.get("src", ""))
        elif name == "div":
            for child in node.find_all(["h1","h2","h3","h4","p","ul","ol","table"], recursive=True):
                pass

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    md_path, out_base, title, subtitle, badges_csv, date_str = sys.argv[1:7]
    badges = badges_csv.split("|")
    html_path = out_base + ".html"
    pdf_path = out_base + ".pdf"
    docx_path = out_base + ".docx"
    build_html(md_path, html_path, title, subtitle, badges, date_str)
    ok_pdf = html_to_pdf(html_path, pdf_path)
    html_to_docx(html_path, docx_path, title)
    print("PDF:", "OK" if ok_pdf else "FAIL", pdf_path)
    print("DOCX: OK", docx_path)
